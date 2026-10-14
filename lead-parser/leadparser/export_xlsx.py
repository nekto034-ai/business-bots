"""Выгрузка лидов в Excel.

Два листа — «С amoCRM» и «Без amoCRM», чтобы сравнить конверсию звонков,
плюс лист «Конверсия», который сам считает результаты по обоим листам.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = [
    ("Балл", 7),
    ("Почему подходит", 60),
    ("Компания", 34),
    ("Чем занимаются", 28),
    ("Телефоны", 20),
    ("Результат звонка", 18),
    ("Комментарий", 30),
    ("Время (от МСК)", 10),
    ("Город", 14),
    ("CRM", 30),
    ("Откуда нашли", 22),
    ("Сайт", 30),
    ("Email", 26),
    ("Соцсети / мессенджеры", 36),
    ("Отзывов в 2ГИС", 9),
    ("Рейтинг", 8),
    ("В 2ГИС с", 11),
    ("Вакансии hh", 40),
    ("Адрес", 34),
    ("Карточка 2ГИС", 16),
]
COL = {name: i for i, (name, _) in enumerate(COLUMNS, 1)}

CALL_RESULTS = ["Не дозвонился", "Не интересно", "Перезвонить", "Интересно", "Встреча", "Сделка"]

SHEETS = [("С amoCRM", True), ("Без amoCRM", False)]

HELP = [
    ("Как читать таблицу", ""),
    ("", ""),
    ("Два листа", "«С amoCRM» — компании, у которых нашли следы amoCRM. «Без amoCRM» — все остальные. Звони по обоим — на листе «Конверсия» увидишь, где результат лучше."),
    ("Балл", "Чем выше, тем больше компания похожа на твоего клиента. Списки отсортированы по баллу — звони сверху вниз."),
    ("Почему подходит", "Готовый повод для разговора: что именно у компании не так."),
    ("Результат звонка", "Выбирай из списка в ячейке. Лист «Конверсия» посчитает итоги сам."),
    ("Время (от МСК)", "Разница с Москвой. «МСК» — то же время, что у тебя."),
    ("CRM", "Какая CRM у компании и откуда это известно: из вакансии на hh или по форме/чату на сайте. Пусто — не нашли следов, это не значит, что CRM нет."),
    ("Откуда нашли", "«hh: вакансии с amoCRM» — компания сама пишет про amoCRM в вакансиях. Иначе — поисковый запрос в 2ГИС."),
    ("Вакансии hh", "Если компания ищет менеджера по продажам — у неё точно есть продажи, а значит и заявки, которые надо где-то вести."),
    ("", ""),
    ("Откуда данные", "2ГИС (открытые карточки компаний), сайты самих компаний, hh.ru (открытые вакансии)."),
    ("Фильтры", "Только компании без филиалов, не больше 15 отзывов в 2ГИС, с телефоном; без федеральных сетей и без конкурентов (интеграторы CRM, веб-студии, разработчики ПО). Настраивается в config.toml."),
]

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
TOP_FILL = PatternFill("solid", fgColor="E2EFDA")
LINK_FONT = Font(color="0563C1", underline="single")


def has_amocrm(lead):
    return bool(lead.get("amocrm") or "amoCRM" in (lead["site"].get("crm") or []))


def _fmt_offset(h):
    if h is None:
        return ""
    return "МСК" if h == 0 else f"{h:+d}"


def _header(ws):
    for i, (title, width) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30


def _row(lead):
    hh_txt = "\n".join(f"{v['title']} — {v['url']}" for v in (lead["hh"] or [])[:3])
    return {
        "Балл": lead["score"],
        "Почему подходит": "; ".join(lead["why"]),
        "Компания": lead["name"],
        "Чем занимаются": lead["rubric"],
        "Телефоны": "\n".join(lead["contacts"]["phones"]),
        "Время (от МСК)": _fmt_offset(lead["msk_offset"]),
        "Город": lead["city"],
        "CRM": lead.get("crm", ""),
        "Откуда нашли": lead["query"],
        "Сайт": lead["site"]["url"] or "",
        "Email": "\n".join(lead["contacts"]["emails"]),
        "Соцсети / мессенджеры": "\n".join(lead["contacts"]["socials"]),
        "Отзывов в 2ГИС": lead["reviews"],
        "Рейтинг": lead["rating"] or "",
        "В 2ГИС с": lead["created"] or "",
        "Вакансии hh": hh_txt,
        "Адрес": lead["address"],
        "Карточка 2ГИС": "открыть",
    }


def _leads_sheet(wb, title, leads):
    ws = wb.create_sheet(title)
    _header(ws)
    last = max(2, len(leads) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(CALL_RESULTS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    res_col = get_column_letter(COL["Результат звонка"])
    dv.add(f"{res_col}2:{res_col}{max(last, 500)}")

    for r, lead in enumerate(leads, 2):
        for name, val in _row(lead).items():
            cell = ws.cell(row=r, column=COL[name], value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        link = ws.cell(row=r, column=COL["Карточка 2ГИС"])
        link.hyperlink = lead["url"]
        link.font = LINK_FONT
        site = lead["site"]["url"]
        if site:
            c = ws.cell(row=r, column=COL["Сайт"])
            c.hyperlink = site if "://" in site else "http://" + site
            c.font = LINK_FONT
        if r <= 11:  # первая десятка — подсветка
            ws.cell(row=r, column=1).fill = TOP_FILL

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"


def _conversion_sheet(wb, counts):
    ws = wb.create_sheet("Конверсия")
    res_col = get_column_letter(COL["Результат звонка"])
    ws.column_dimensions["A"].width = 22
    for i, (title, _) in enumerate(SHEETS):
        ws.column_dimensions[get_column_letter(i + 2)].width = 16
    ws.cell(row=1, column=1, value="Результат звонка").font = Font(bold=True)
    for i, (title, _) in enumerate(SHEETS):
        c = ws.cell(row=1, column=i + 2, value=title)
        c.font = Font(bold=True)

    def rng(title):
        return f"'{title}'!{res_col}2:{res_col}1000"

    ws.cell(row=2, column=1, value="Компаний в списке")
    for i, (title, _) in enumerate(SHEETS):
        ws.cell(row=2, column=i + 2, value=counts[title])
    ws.cell(row=3, column=1, value="Обзвонено")
    for i, (title, _) in enumerate(SHEETS):
        ws.cell(row=3, column=i + 2, value=f"=COUNTA({rng(title)})")
    for j, res in enumerate(CALL_RESULTS):
        row = 4 + j
        ws.cell(row=row, column=1, value=res)
        for i, (title, _) in enumerate(SHEETS):
            ws.cell(row=row, column=i + 2, value=f'=COUNTIF({rng(title)},"{res}")')
    row = 4 + len(CALL_RESULTS) + 1
    ws.cell(row=row, column=1, value="Конверсия в интерес, %").font = Font(bold=True)
    for i, (title, _) in enumerate(SHEETS):
        col = get_column_letter(i + 2)
        first = 4 + CALL_RESULTS.index("Интересно")
        ws.cell(row=row, column=i + 2,
                value=f"=IF({col}3=0,\"—\",ROUND(SUM({col}{first}:{col}{3 + len(CALL_RESULTS)})/{col}3*100,0))").font = Font(bold=True)
    ws.cell(row=row + 1, column=1,
            value="(Интересно + Встреча + Сделка) / Обзвонено").font = Font(italic=True, color="808080")


def write(leads, path):
    wb = Workbook()
    wb.remove(wb.active)
    counts = {}
    for title, amo in SHEETS:
        part = [L for L in leads if has_amocrm(L) == amo]
        counts[title] = len(part)
        _leads_sheet(wb, title, part)
    _conversion_sheet(wb, counts)

    h = wb.create_sheet("Как читать")
    h.column_dimensions["A"].width = 20
    h.column_dimensions["B"].width = 100
    for r, (a, b) in enumerate(HELP, 1):
        h.cell(row=r, column=1, value=a).font = Font(bold=True, size=14 if r == 1 else 11)
        h.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
